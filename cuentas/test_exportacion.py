from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
import gestor as g
from exportacion import filtered_excel, cut_excel
import cortes
import repartos
from test_revision_masiva import seed


class ExportTests(unittest.TestCase):
    def test_legacy_cut_adjustments_keep_signed_fractional_amounts(self):
        with tempfile.TemporaryDirectory() as folder:
            db=g.connect(Path(folder)/'test.db')
            try:
                seed(db,1)
                with db:db.execute("UPDATE movements SET state='aceptado'")
                filters=dict(start='2026-09-01',end='2026-09-30')
                plan=cortes.plan(db,filters)
                cortes.save(db,dict(filters,token=plan['token']))
                g.review(db,1,'editar',mi='499.5',amor='500.5')
                plan=cortes.plan(db,filters)
                cid=cortes.save(db,dict(filters,token=plan['token']))['id']
                wb=load_workbook(BytesIO(cut_excel(db,cid)))
                ws=wb['Detalle']
                self.assertEqual(ws['F2'].value,True)
                self.assertEqual(ws['G2'].value,0)
                self.assertEqual(ws['H2'].value,-0.5)
                self.assertEqual(ws['I2'].value,0.5)
                wb.close()
            finally:db.close()

    def test_cut_exports_saved_values_and_names_not_current_ledger(self):
        with tempfile.TemporaryDirectory() as folder:
            db=g.connect(Path(folder)/'test.db')
            try:
                seed(db,2)
                repartos.migrate(db)
                repartos.sync(db,[dict(id='a',name='Ana',surname='',alias='',active=1,position=0),
                                   dict(id='b',name='Bea',surname='',alias='',active=1,position=1)])
                with db:db.execute("UPDATE movements SET state='aceptado',description='=1+1'")
                filters=dict(start='2026-09-01',end='2026-09-30')
                plan=cortes.plan(db,filters)
                cid=cortes.save(db,dict(filters,token=plan['token'],note='=malicious'))['id']
                with db:
                    g.categorias.ensure(db,'Otro')
                    db.execute("UPDATE movements SET description='CAMBIADO',category='Otro'")
                    db.execute("UPDATE household_people SET name='Nuevo'")
                    db.execute('UPDATE allocations SET amount=0')
                cortes.cancel(db,cid)
                wb=load_workbook(BytesIO(cut_excel(db,cid)))
                ws=wb['Detalle']
                self.assertEqual(ws.max_row,3)
                self.assertEqual(ws['D2'].value,'=1+1')
                self.assertEqual(ws['D2'].data_type,'s')
                self.assertEqual(ws['E2'].value,'Comida')
                self.assertEqual(ws['H1'].value,'Ana [a] CLP')
                self.assertEqual(ws['H2'].value,500)
                self.assertEqual(ws['G2'].value,1000)
                self.assertEqual(wb['Corte']['B5'].value,'Anulado')
                self.assertEqual(wb['Corte']['B6'].data_type,'s')
                wb.close()
                with self.assertRaises(ValueError):cut_excel(db,999)
            finally:db.close()

    def test_filters_all_pages_and_cell_types(self):
        with tempfile.TemporaryDirectory() as folder:
            db=g.connect(Path(folder)/'test.db')
            try:
                seed(db,80)
                with db:
                    db.execute("UPDATE movements SET period='2025-09' WHERE id>40")
                    db.execute("UPDATE movements SET description='=1+1',mi=-500000,amor=1000500000 WHERE id=1")
                wb=load_workbook(BytesIO(filtered_excel(db,dict(year='2026',month='09',category='Comida',page='2'))))
                ws=wb.active
                self.assertEqual(ws.max_row,41)
                self.assertEqual([c.value for c in ws[1]],[c[0] for c in db.execute('SELECT * FROM v_powerbi LIMIT 0').description])
                row=next(r for r in ws.iter_rows(min_row=2) if r[0].value==1)
                self.assertEqual(row[3].value,'=1+1')
                self.assertEqual(row[3].data_type,'s')
                self.assertEqual(row[6].value,-0.5)
                self.assertEqual(row[1].value.strftime('%Y-%m-%d'),'2026-09-01')
                self.assertEqual(ws.freeze_panes,'A2')
                wb.close()
                wb=load_workbook(BytesIO(filtered_excel(db,dict(year='2030'))))
                self.assertEqual(wb.active.max_row,1)
                wb.close()
            finally:db.close()
