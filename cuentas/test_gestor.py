import tempfile
import unittest
from datetime import datetime, date
from pathlib import Path
from unittest.mock import patch

import gestor as g


def item(desc='COMPRA TEST',amount=1000,dt='2026-09-15',p='2026-09',kind='gasto'):
    a=g.money(amount)
    return dict(sheet='Movimientos',row=2,raw={'descripcion':desc},purchase_date=dt,
                period=p,description=desc,normalized_description=g.norm(desc),kind=kind,
                amount=a,purchase_amount=a,installments_pending=0,installment=0,
                holder='Titular',mi=a//2,amor=a-a//2)


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.folder=Path(self.temp.name)
        self.path=self.folder/'db.sqlite3'
        self.db=g.connect(self.path)

    def tearDown(self):
        self.db.close();self.temp.cleanup()

    def load(self,name,items):
        source=self.folder/name;source.write_bytes(name.encode())
        with patch.object(g,'bank_rows',return_value=items):
            return g.import_bank(self.db,self.path,source,'2026-09')

    def test_identical_file_does_not_repeat_or_reset_decisions(self):
        result=self.load('a.xlsx',[item()])
        g.review(self.db,1,'aceptar',category='Comida')
        again=self.load('a.xlsx',[item()])
        self.assertTrue(again['reutilizado'])
        self.assertEqual(result['lote'],again['lote'])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM movements').fetchone()[0],1)
        self.assertEqual(g.summary(self.db)['real_confirmado_clp'],'1000')
        self.assertEqual(self.db.execute('SELECT category FROM movements').fetchone()[0],'Comida')

    def test_soft_delete_is_audited_hidden_and_recognized_on_import(self):
        import revision_masiva as bulk
        self.load('a.xlsx',[item()])
        g.review(self.db,1,'descartar',note='No considerar')
        self.assertEqual(g.summary(self.db)['bruto_clp'],'0')
        clause,args=bulk.where({})
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM movements m'+clause,args).fetchone()[0],0)
        clause,args=bulk.where({'state':'excluido_gasto'})
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM movements m'+clause,args).fetchone()[0],1)
        self.load('b.xlsx',[item(desc='COMPRA TEST*')])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM movements').fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT state FROM movements').fetchone()[0],'excluido_gasto')
        g.review(self.db,1,'restaurar')
        self.assertEqual(self.db.execute('SELECT state FROM movements').fetchone()[0],'pendiente')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM decisions').fetchone()[0],2)

    def test_soft_delete_protects_active_cuts(self):
        import cortes
        cortes.migrate(self.db)
        self.load('a.xlsx',[item()])
        with self.db:
            self.db.execute("INSERT INTO cuts(id,created_at,start_date,end_date,note,token) VALUES(1,'now','2026-09-01','2026-09-30','','test')")
            self.db.execute("INSERT INTO cut_items(cut_id,movement_id,mi,amor,snapshot) VALUES(1,1,0,0,'{}')")
        with self.assertRaisesRegex(ValueError,'corte vigente'):
            g.review(self.db,1,'descartar')
        self.assertEqual(self.db.execute('SELECT state FROM movements').fetchone()[0],'pendiente')

    def test_two_real_identical_purchases_are_preserved(self):
        a=item();b=item();b['row']=3
        self.load('a.xlsx',[a,b])
        self.assertEqual(g.summary(self.db)['bruto_clp'],'2000')
        self.assertEqual(g.summary(self.db)['deduplicado_propuesto_clp'],'1000')
        g.review(self.db,1,'aceptar',category='Comida')
        g.review(self.db,2,'aceptar',category='Comida')
        self.assertEqual(g.summary(self.db)['real_confirmado_clp'],'2000')

    def test_overlapping_file_preserves_validation(self):
        self.load('a.xlsx',[item()]);g.review(self.db,1,'aceptar',category='Comida')
        self.load('b.xlsx',[item()])
        self.assertEqual(g.summary(self.db)['real_confirmado_clp'],'1000')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM movements').fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM observations').fetchone()[0],2)
        self.assertEqual(g.summary(self.db,2)['real_confirmado_clp'],'1000')

    def test_overlap_multiplicity_and_reordered_rows(self):
        a=item();b=item();b['row']=3
        self.load('a.xlsx',[a,b])
        g.review(self.db,1,'aceptar',category='Comida')
        g.review(self.db,2,'aceptar',category='Comida')
        result=self.load('b.xlsx',[b,a])
        self.assertEqual(result['reconocidos'],2)
        c=item();c['row']=4
        self.assertEqual(self.load('c.xlsx',[a,b,c])['reconocidos'],0)

    def test_overlap_preserves_edits_and_only_new_pending(self):
        self.load('a.xlsx',[item()])
        g.review(self.db,1,'aceptar',category='Comida',mi='300',amor='700')
        new=item('COMPRA NUEVA');new['row']=3
        result=self.load('b.xlsx',[item(),new])
        self.assertEqual((result['reconocidos'],result['nuevos']),(1,1))
        self.assertEqual(tuple(self.db.execute('SELECT mi,amor,state FROM movements WHERE id=1').fetchone()),(g.money(300),g.money(700),'aceptado'))

    def test_changed_bank_name_reuses_original(self):
        self.load('old.xlsx',[item('COMPRA SUPERMERCADO DE EJEMPLO*')])
        g.review(self.db,1,'aceptar',category='Supermercado',mi='300',amor='700')
        before=dict(self.db.execute('SELECT * FROM movements WHERE id=1').fetchone())
        result=self.load('new.xlsx',[item('COMPRA SUPERMERCADO DE EJEMPLO 2')])
        self.assertEqual((result['reconocidos'],result['nuevos']),(1,0))
        self.assertEqual(dict(self.db.execute('SELECT * FROM movements WHERE id=1').fetchone()),before)
        self.assertEqual(self.db.execute('SELECT count(*) FROM observations').fetchone()[0],2)

    def test_similar_multiple_purchases_require_review(self):
        self.load('old.xlsx',[item('COMPRA SUPERMERCADO DE EJEMPLO*')])
        a=item('COMPRA SUPERMERCADO DE EJEMPLO 2');b=item('COMPRA SUPERMERCADO DE EJEMPLO 3');b['row']=3
        result=self.load('new.xlsx',[a,b])
        self.assertEqual(result['reconocidos'],0)
        self.assertGreater(self.db.execute('SELECT count(*) FROM candidates').fetchone()[0],0)

    def test_exact_purchase_cannot_also_match_similar_purchase(self):
        a=item('COMPRA SUPERMERCADO DE EJEMPLO')
        self.load('old.xlsx',[a])
        b=item('COMPRA SUPERMERCADO DE EJEMPLO 2');b['row']=3
        result=self.load('new.xlsx',[a,b])
        self.assertEqual((result['reconocidos'],result['nuevos']),(1,1))

    def test_changed_date_amount_holder_or_installment_not_reused(self):
        self.load('old.xlsx',[item('COMPRA SUPERMERCADO DE EJEMPLO*')])
        for index,(field,value) in enumerate([('purchase_date','2026-09-16'),('amount',g.money(1001)),
                     ('holder','Adicional'),('period','2026-10'),('purchase_amount',g.money(9000))]):
            row=item('COMPRA SUPERMERCADO DE EJEMPLO 2');row[field]=value
            self.assertEqual(self.load(f'new{index}.xlsx',[row])['reconocidos'],0)

    def test_archived_excel_datetime_overlap(self):
        original=item(dt='2026-09-19')
        original['raw']={'fecha':datetime(2026,9,19), 'descripcion':'COMPRA TEST',
                         'monto':1000, 'valor cuota':1000, 'cuotas pendientes':0,
                         'titular/adicional':'Titular'}
        self.load('old.xlsx',[original])
        g.review(self.db,1,'aceptar',category='Comida',mi='300',amor='700')
        before=dict(self.db.execute('SELECT * FROM movements WHERE id=1').fetchone())
        new=item('COMPRA NUEVA');new['row']=3
        result=self.load('new.xlsx',[original,new])
        self.assertEqual((result['reconocidos'],result['nuevos']),(1,1))
        self.assertEqual(dict(self.db.execute('SELECT * FROM movements WHERE id=1').fetchone()),before)
        self.assertEqual(self.db.execute('SELECT count(*) FROM movements').fetchone()[0],2)

    def test_refunds_and_card_payment(self):
        a=item(amount=-100,kind='devolucion');b=item('PAGO TARJETA CMR',-500,kind='pago_tarjeta');b['row']=3
        self.load('a.xlsx',[a,b]);g.review(self.db,1,'aceptar',category='Retail')
        self.assertEqual(g.summary(self.db)['real_confirmado_clp'],'-100')
        self.assertEqual(g.summary(self.db)['bruto_clp'],'-100')
        with self.assertRaises(ValueError):g.review(self.db,2,'aceptar',category='Pago')

    def test_installments_in_distinct_periods_not_duplicates(self):
        a=item(dt='2026-05-10',p='2026-08');b=item(dt='2026-05-10',p='2026-09')
        self.load('a.xlsx',[a]);self.load('b.xlsx',[b])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidates').fetchone()[0],0)

    def test_invalid_split_and_atomic_import(self):
        self.load('a.xlsx',[item()])
        with self.assertRaises(ValueError):g.review(self.db,1,'aceptar',category='Comida',mi='0',amor='0')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM decisions').fetchone()[0],0)
        bad=item();bad['kind']='bad'
        with self.assertRaises(Exception):self.load('b.xlsx',[bad])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM batches').fetchone()[0],1)

    def test_exact_half_peso_and_float_residue(self):
        self.assertEqual(g.money('59996.79999999999'),59996800000)
        self.assertEqual(g.money('5527.5')*2,g.money('11055'))
        self.assertIsNone(g.money('-',True))

    def test_rename_does_not_duplicate_and_period_cannot_silently_change(self):
        self.load('a.xlsx',[item()])
        renamed=self.folder/'renamed.xlsx';renamed.write_bytes(b'a.xlsx')
        self.assertTrue(g.import_bank(self.db,self.path,renamed,'2026-09')['reutilizado'])
        with self.assertRaises(ValueError):g.import_bank(self.db,self.path,renamed,'2026-10')

    def test_preserve_real_purchase_does_not_approve_category(self):
        a=item();b=item();b['row']=3
        self.load('a.xlsx',[a,b])
        g.review(self.db,2,'conservar',note='Dos compras reales')
        self.assertEqual(self.db.execute('SELECT state FROM movements WHERE id=2').fetchone()[0],'pendiente')
        self.assertEqual(self.db.execute("SELECT resolved FROM issues WHERE movement_id=2 AND code='posible_duplicado'").fetchone()[0],1)
        self.assertEqual(g.summary(self.db)['real_confirmado_clp'],'0')

    def test_save_partial_review_and_refresh_issues(self):
        self.load('a.xlsx',[item()])
        g.review(self.db,1,'editar',category='Comida',mi='',amor='')
        row=self.db.execute('SELECT * FROM movements WHERE id=1').fetchone()
        self.assertEqual(row['category'],'Comida')
        self.assertIsNone(row['mi'])
        codes=[r[0] for r in self.db.execute('SELECT code FROM issues WHERE movement_id=1 AND resolved=0')]
        self.assertNotIn('categoria',codes)
        self.assertIn('reparto',codes)
        with self.assertRaises(ValueError):g.review(self.db,1,'aceptar')
        g.review(self.db,1,'editar',mi='500',amor='500')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM issues WHERE movement_id=1 AND resolved=0').fetchone()[0],0)

    def test_special_case_persists_and_allows_unbalanced_confirmation(self):
        import revision_masiva as bulk
        self.load('a.xlsx',[item()])
        g.review(self.db,1,'editar',category='Comida',mi='0',amor='0',special_case=True)
        info,_=bulk.plan(self.db,{})
        self.assertEqual(info['eligible'],1)
        bulk.confirm(self.db,{},info['token'])
        g.review(self.db,1,'editar',mi='100',amor='200')
        r=self.db.execute('SELECT * FROM movements').fetchone()
        self.assertEqual((r['special_case'],r['state'],r['amount']),(1,'aceptado',g.money(1000)))
        with self.assertRaises(ValueError):g.review(self.db,1,'editar',special_case=False)
        with self.assertRaises(ValueError):g.review(self.db,1,'editar',mi='',special_case=True)
        g.review(self.db,1,'editar',mi='400',amor='600',special_case=False)
        self.assertEqual(self.db.execute('SELECT caso_especial FROM v_powerbi').fetchone()[0],0)

    def test_special_case_individual_and_boolean_validation(self):
        self.load('a.xlsx',[item()])
        with self.assertRaises(ValueError):g.review(self.db,1,'aceptar',category='Comida',mi='0',amor='0')
        g.review(self.db,1,'aceptar',category='Comida',mi='0',amor='0',special_case=True)
        with self.assertRaises(ValueError):g.review(self.db,1,'editar',special_case='false')
        self.assertEqual(self.db.execute('SELECT caso_especial FROM v_gastos_confirmados').fetchone()[0],1)

    def test_export_view_only_counts_accepted(self):
        self.load('a.xlsx',[item()])
        g.export(self.db,self.folder/'out')
        self.assertEqual(len((self.folder/'out/v_gastos_confirmados.csv').read_text(encoding='utf-8-sig').splitlines()),1)
        g.review(self.db,1,'aceptar',category='Comida')
        g.export(self.db,self.folder/'out')
        self.assertEqual(len((self.folder/'out/v_gastos_confirmados.csv').read_text(encoding='utf-8-sig').splitlines()),2)


class SourceTests(unittest.TestCase):
    def test_dates_and_serialized_datetimes(self):
        for value in (date(2026,9,19),datetime(2026,9,19,13,25),
                      '2026-09-19','19/09/2026','19-09-2026',
                      '2026-09-19 00:00:00','2026-09-19T13:25:59',
                      '2026-09-19 00:00:00.123456','2026-09-19T00:00:00.123456'):
            with self.subTest(value=value):self.assertEqual(g.day(value),'2026-09-19')
        for value in ('2026-02-30 00:00:00','2026-09-19 garbage','2026-09-19 25:00:00'):
            with self.subTest(value=value),self.assertRaises(ValueError):g.day(value)

    def test_foreign_currency_and_bad_dates_rejected(self):
        class Sheet:
            title='Movimientos'
            def __init__(self,values):self.values=values
            def iter_rows(self,values_only=True):return iter(self.values)
        class Book(list):
            def close(self):pass
        headers=['FECHA','DESCRIPCION','TITULAR/ADICIONAL','MONTO','CUOTAS PENDIENTES','VALOR CUOTA','MONEDA']
        for row,expected in [(['2026-09-01','Compra','Titular',10,0,10,'USD'],'Moneda extranjera'),(['sin fecha','Compra','Titular',10,0,10,'CLP'],'Fecha no interpretable')]:
            with patch('openpyxl.load_workbook',return_value=Book([Sheet([headers,row])])):
                with self.assertRaisesRegex(ValueError,expected):g.bank_rows(b'fixture','2026-09')

    def test_synthetic_bank_export(self):
        from openpyxl import Workbook
        folder=tempfile.TemporaryDirectory();self.addCleanup(folder.cleanup)
        source=Path(folder.name)/'synthetic.xlsx'
        book=Workbook();sheet=book.active
        sheet.append(['FECHA','DESCRIPCION','TITULAR/ADICIONAL','MONTO','CUOTAS PENDIENTES','VALOR CUOTA'])
        sheet.append([datetime(2026,5,1),'COMPRA EJEMPLO','Titular',30000,2,10000])
        sheet.append([datetime(2026,9,1),'PAGO TARJETA CMR','Titular',20000,0,-20000])
        for _ in range(2):sheet.append([datetime(2026,9,2),'COMPRA PRUEBA','Titular',1500,0,1500])
        book.save(source);book.close()
        with self.assertRaisesRegex(ValueError,'mes de las cuotas'):
            g.bank_rows(source.read_bytes(),None)
        rows=g.bank_rows(source.read_bytes(),'2026-09')
        self.assertEqual(len(rows),4)
        self.assertEqual(sum(r['kind']=='pago_tarjeta' for r in rows),1)
        installment=[r for r in rows if r['installment']]
        self.assertEqual(len(installment),1)
        self.assertEqual(installment[0]['purchase_date'],'2026-05-01')
        self.assertEqual(installment[0]['period'],'2026-09')
        self.assertEqual(installment[0]['amount'],g.money(10000))
        self.assertEqual(sum(r['description']=='COMPRA PRUEBA' and r['amount']==g.money(1500) for r in rows),2)


if __name__=='__main__':unittest.main()
