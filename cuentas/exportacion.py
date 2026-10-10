"""Excel downloads for the exact ledger filter, without pagination."""
from datetime import date
from io import BytesIO
import io
import csv
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import revision_masiva as bulk
import repartos
import cortes


def cut_excel(db, cid):
    data=cortes.detail(db,cid)
    cut=data['cut']; rows=data['rows']
    names={}
    for row in rows:
        for pid in row.get('allocations',{}):
            names.setdefault(pid,row.get('people',{}).get(pid,pid))
    wb=Workbook(); summary=wb.active; summary.title='Corte'
    for record in [('Corte',cid),('Desde',cut['start_date']),('Hasta',cut['end_date']),
                   ('Validado',cut['created_at']),('Estado','Anulado' if cut['cancelled_at'] else 'Vigente'),
                   ('Nota',cut['note']),('Movimientos',len(rows))]:
        summary.append(record)
    summary.append(['Persona','Total CLP'])
    totals=data['allocations'] if names else {'Mi':data['mi'],'Amor':data['amor']}
    for pid, amount in totals.items():summary.append([names.get(pid,pid),Decimal(amount)/1_000_000])
    ws=wb.create_sheet('Detalle')
    headers=['ID','Fecha compra','Periodo','Descripcion','Categoria al validar','Ajuste','Total del corte CLP']
    headers += [f'{name} [{pid}] CLP' for pid,name in names.items()] if names else ['Mi CLP','Amor CLP']
    ws.append(headers)
    for row in rows:
        allocated=row.get('allocations',{})
        amounts=[allocated.get(pid,0) or 0 for pid in names] if names else [row['mi'],row['amor']]
        ws.append([row['id'],date.fromisoformat(row['purchase_date']),row['period'],row['description'],
                   row['category'] or '',bool(row.get('adjustment')),Decimal(sum(amounts))/1_000_000]+
                  [Decimal(n)/1_000_000 for n in amounts])
    for sheet in wb:
        sheet.freeze_panes='A2'; sheet.sheet_view.showGridLines=False
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value,str):cell.data_type='s'
                elif isinstance(cell.value,date):cell.number_format='dd-mm-yyyy'
                elif isinstance(cell.value,Decimal):cell.number_format='#,##0.######;[Red]-#,##0.######'
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='126B57')
        for i in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(i)].width=24
    summary.column_dimensions['B'].width=65
    summary['B6'].alignment=Alignment(wrap_text=True)
    ws.column_dimensions['D'].width=55;ws.auto_filter.ref=ws.dimensions
    output=BytesIO();wb.save(output);wb.close()
    return output.getvalue()


def records(db, filters):
    clause,args=bulk.where(filters)
    cursor=db.execute('SELECT v.* FROM movements m JOIN v_powerbi v ON v.id=m.id'+clause+' ORDER BY m.purchase_date DESC,m.id DESC',args)
    members=repartos.people(db)
    headers=[c[0] for c in cursor.description if not members or c[0] not in ('mi_clp','amor_clp')]
    rows=[]
    for record in cursor:
        row=[record[h] for h in headers]
        values=repartos.get(db,record['id'])
        if members:row.extend(None if values.get(p['id'],0) is None else Decimal(values.get(p['id'],0))/1_000_000 for p in members)
        rows.append(row)
    headers.extend(f"{p['name']} [{p['id']}] CLP" for p in members)
    return headers,rows


def household_csv(db, confirmed):
    filters=dict(state='aceptado') if confirmed else dict(show_duplicates='1')
    headers,rows=records(db,filters)
    if confirmed:
        kind=headers.index('tipo')
        rows=[r for r in rows if r[kind] in ('gasto','devolucion')]
    out=io.StringIO();writer=csv.writer(out)
    for row in [headers]+rows:
        writer.writerow(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
    return out.getvalue().encode('utf-8-sig')


def filtered_excel(db, filters):
    headers,cursor=records(db,filters)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Movimientos'
    ws.append(headers)
    for record in cursor:
        ws.append(list(record))
        for cell, header in zip(ws[ws.max_row], headers):
            if isinstance(cell.value, str):
                # Treat user-entered text as literal text, never as an Excel formula.
                cell.data_type = 's'
            if header == 'fecha_compra' and cell.value:
                cell.value = date.fromisoformat(cell.value)
                cell.number_format = 'dd-mm-yyyy'
            elif header in ('monto_clp','mi_clp','amor_clp') or header.endswith(' CLP'):
                cell.number_format = '#,##0.######;[Red]-#,##0.######'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor='126B57')
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 25
    for index, header in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(index)].width = (
            55 if header=='descripcion' else 24 if header=='categoria' else 19)
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()
